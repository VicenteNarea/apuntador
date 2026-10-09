"""Lectura de documentos: Markdown limpio, PowerPoint con grupos y prioridad al documento nombrado."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import asistente_llamadas as A  # noqa: E402


def test_limpiar_md():
    assert A.limpiar_md("- **MLP:** red `neuronal`") == "- MLP: red neuronal"
    assert A.limpiar_md("> **Objetivo técnico:** transformar") == "Objetivo técnico: transformar"
    assert A.limpiar_md("---") == ""
    assert A.limpiar_md("|---|:---:|") == ""
    assert A.limpiar_md("| Etapa | Técnica |") == "Etapa · Técnica"
    assert A.limpiar_md("ver [el informe](http://x.y/z.pdf)") == "ver el informe"


def test_pptx_lee_cuadros_agrupados(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text = "Estudio económico"
    grupo = sl.shapes.add_group_shape()
    grupo.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = "Ahorro USD 630 al mes"
    sl.notes_slide.notes_text_frame.text = "Decir que el VAN es positivo"
    ruta = tmp_path / "presentacion.pptx"
    prs.save(ruta)
    (titulo, texto), = A.leer_secciones(ruta)
    assert titulo.startswith("Diapositiva 1: Estudio económico")
    assert "Ahorro USD 630 al mes" in texto and "Guion: Decir que el VAN es positivo" in texto


def _kb(tmp_path):
    (tmp_path / "Presentacion_Defensa.md").write_text(
        "# Defensa\n\n## Diapositiva 6 · Justificación\n\nAhorro de **USD 630** mensuales.\n", encoding="utf-8")
    (tmp_path / "Informe_Final.md").write_text(
        "# Informe\n\n## Justificación del proyecto\n\nLa justificación es reducir 40 horas al mes de trabajo manual.\n",
        encoding="utf-8")
    (tmp_path / "Guia.md").write_text(
        "# Guía\n\n## Justificación\n\nJustificación: trazabilidad, justificación auditable, justificación BHP.\n",
        encoding="utf-8")
    kb = A.BaseConocimiento(tmp_path)
    kb.cargar()
    return kb


def test_prioriza_el_documento_nombrado(tmp_path):
    kb = _kb(tmp_path)
    assert kb.buscar("¿Qué dice la presentación sobre la justificación?")[0][0] == "Presentacion_Defensa.md"
    assert kb.buscar("¿Qué dice la diapo de la justificación?")[0][0] == "Presentacion_Defensa.md"
    assert kb.buscar("En el informe, ¿cuál es la justificación?")[0][0] == "Informe_Final.md"


def test_markdown_indexado_sin_marcas(tmp_path):
    kb = _kb(tmp_path)
    texto = [t for f, t in kb.frag if f == "Presentacion_Defensa.md"][0]
    assert "**" not in texto and "USD 630" in texto and texto.startswith("[Diapositiva 6 · Justificación]")
